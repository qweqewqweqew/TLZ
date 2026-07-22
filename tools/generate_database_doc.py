from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"D:\QtProj\tongliiz\MzTLZ\MzTLZ数据库表结构说明.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D0D7DE")


def set_table_geometry(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    total_dxa = int(sum(widths) * 1440)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_dxa))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "0")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)


def set_cell_width(cell, width):
    dxa = str(int(width * 1440))
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), dxa)
    cell.width = Inches(width)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    if widths:
        set_table_geometry(table, widths)
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        if widths:
            set_cell_width(hdr[i], widths[i])
        set_cell_shading(hdr[i], "E8EEF5")
        set_cell_text(hdr[i], h, bold=True, color="0B2545", size=8.5)
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if widths:
                set_cell_width(cells[i], widths[i])
            set_cell_text(cells[i], value, size=8.5)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = "Heading {}".format(level)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.color.rgb = RGBColor(46, 116, 181)
    return p


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11)
section.page_height = Inches(8.5)
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.6)
section.right_margin = Inches(0.6)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10)
for style_name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(46, 116, 181)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("MzTLZ 数据库表结构说明")
r.bold = True
r.font.name = "Microsoft YaHei"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(20)
r.font.color.rgb = RGBColor(11, 37, 69)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("整理时间：2026-07-17    数据源：localhost\\SQLEXPRESS / MzTLZ")
r.font.name = "Microsoft YaHei"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(85, 85, 85)

add_heading(doc, "1. 数据库概览", 1)
add_table(
    doc,
    ["项目", "内容"],
    [
        ["数据库类型", "Microsoft SQL Server 2022 Express"],
        ["连接方式", "Qt QODBC"],
        ["服务器", "localhost\\SQLEXPRESS"],
        ["数据库", "MzTLZ"],
        ["认证方式", "Windows 集成认证（Trusted_Connection=Yes）"],
        ["主要访问代码", "Database.cpp / InspectionRepository.cpp"],
    ],
    [2.2, 7.3],
)

add_heading(doc, "2. 当前业务表与数据量", 1)
add_table(
    doc,
    ["表名", "用途", "当前行数"],
    [
        ["dbo.ProcessParameter", "当前正在使用的工艺配方/路径参数", "6"],
        ["dbo.InspectionRecord", "一次加工/检测任务的主记录", "5"],
        ["dbo.ParticleDetection", "一次检测中识别到的颗粒明细", "41"],
        ["dbo.ProcessParameterHistory", "任务开始时冻结的工艺参数快照", "0"],
    ],
    [2.5, 5.7, 1.3],
)

field_tables = {
    "dbo.ProcessParameter": [
        ("Id", "INT", "否", "主键，自增"),
        ("RecipeName", "NVARCHAR(200)", "否", "配方名"),
        ("PathIndex", "INT", "否", "路径序号"),
        ("MaxParticleHeight", "FLOAT", "是", "最大颗粒高度/工艺允许上限"),
        ("TargetX", "FLOAT", "是", "目标坐标 X"),
        ("TargetY", "FLOAT", "是", "目标坐标 Y"),
        ("TargetZ", "FLOAT", "是", "目标坐标 Z"),
        ("FeedSpeed", "FLOAT", "是", "进给速度"),
        ("FeedAmount", "FLOAT", "是", "进给量"),
        ("SpindleSpeed", "FLOAT", "是", "主轴转速"),
        ("CutCount", "INT", "是", "切削次数"),
        ("SingleCutAmount", "FLOAT", "是", "单次切削量"),
        ("CreatedAt", "DATETIME2", "否", "创建时间"),
        ("UpdatedAt", "DATETIME2", "否", "更新时间"),
    ],
    "dbo.InspectionRecord": [
        ("Id", "INT", "否", "主键，自增"),
        ("RecipeName", "NVARCHAR(200)", "否", "使用的配方名"),
        ("ProcessStartAt", "DATETIME2", "是", "加工开始时间"),
        ("ProcessEndAt", "DATETIME2", "是", "加工结束时间"),
        ("InspectStartAt", "DATETIME2", "是", "检测开始时间"),
        ("InspectEndAt", "DATETIME2", "是", "检测结束时间"),
        ("ParticleCount", "INT", "否", "检测到的颗粒总数"),
        ("ClearedCount", "INT", "否", "已清除颗粒数"),
        ("MaxParticleHeight", "FLOAT", "是", "本次实测最大颗粒高度"),
        ("OverviewImagePath", "NVARCHAR(1000)", "是", "全景图片路径"),
        ("Remark", "NVARCHAR(1000)", "是", "备注"),
        ("CreatedAt", "DATETIME2", "否", "创建时间"),
        ("UpdatedAt", "DATETIME2", "否", "更新时间"),
    ],
    "dbo.ParticleDetection": [
        ("Id", "BIGINT", "否", "主键，自增"),
        ("InspectionRecordId", "INT", "否", "所属检测记录，关联 InspectionRecord.Id"),
        ("ParticleIndex", "INT", "否", "本次任务内颗粒序号"),
        ("PositionX", "FLOAT", "是", "颗粒位置 X"),
        ("PositionY", "FLOAT", "是", "颗粒位置 Y"),
        ("PositionZ", "FLOAT", "是", "颗粒位置 Z"),
        ("Height", "FLOAT", "是", "颗粒高度"),
        ("IsCleared", "BIT", "否", "是否已清除"),
        ("ClearedAt", "DATETIME2", "是", "清除时间"),
        ("DetectedAt", "DATETIME2", "否", "检测时间"),
    ],
    "dbo.ProcessParameterHistory": [
        ("Id", "BIGINT", "否", "主键，自增"),
        ("InspectionRecordId", "INT", "否", "所属检测记录，关联 InspectionRecord.Id"),
        ("RecipeName", "NVARCHAR(200)", "否", "配方名"),
        ("PathIndex", "INT", "否", "路径序号"),
        ("MaxParticleHeight", "FLOAT", "是", "最大颗粒高度"),
        ("TargetX", "FLOAT", "是", "目标坐标 X"),
        ("TargetY", "FLOAT", "是", "目标坐标 Y"),
        ("TargetZ", "FLOAT", "是", "目标坐标 Z"),
        ("FeedSpeed", "FLOAT", "是", "进给速度"),
        ("FeedAmount", "FLOAT", "是", "进给量"),
        ("SpindleSpeed", "FLOAT", "是", "主轴转速"),
        ("CutCount", "INT", "是", "切削次数"),
        ("SingleCutAmount", "FLOAT", "是", "单次切削量"),
        ("SnapshotAt", "DATETIME2", "否", "快照生成时间"),
    ],
}

add_heading(doc, "3. 字段结构", 1)
for table_name, rows in field_tables.items():
    add_heading(doc, table_name, 2)
    add_table(doc, ["字段", "类型", "可空", "说明"], rows, [2.1, 2.1, 0.8, 4.5])

add_heading(doc, "4. 外键与索引", 1)
add_table(
    doc,
    ["类型", "名称", "表/字段", "说明"],
    [
        ["外键", "FK_ParticleDetection_Inspection", "ParticleDetection.InspectionRecordId", "关联 InspectionRecord.Id"],
        ["外键", "FK_ProcessParameterHistory_Inspection", "ProcessParameterHistory.InspectionRecordId", "关联 InspectionRecord.Id"],
        ["唯一约束", "UQ_ProcessParameterHistory_Plate_Path", "ProcessParameterHistory(InspectionRecordId, PathIndex)", "同一次任务的路径快照不重复"],
        ["索引", "IX_InspectionRecord_RecipeName", "InspectionRecord.RecipeName", "支持按配方筛选"],
        ["索引", "IX_InspectionRecord_InspectStartAt", "InspectionRecord.InspectStartAt", "支持按检测时间查询"],
        ["索引", "IX_ParticleDetection_InspectionRecordId", "ParticleDetection.InspectionRecordId", "支持加载某次任务颗粒明细"],
        ["索引", "IX_ProcessParameterHistory_InspectionRecordId", "ProcessParameterHistory.InspectionRecordId", "支持加载某次任务工艺快照"],
    ],
    [1.1, 3.0, 3.1, 2.3],
)

add_heading(doc, "5. 当前样例数据备注", 1)
add_table(
    doc,
    ["项目", "当前情况"],
    [
        ["ProcessParameter", "DefaultRecipe 的 3 条路径各出现两次，因此共有 6 行。"],
        ["InspectionRecord", "当前 5 条均为模拟任务，备注为“模拟任务 #1”至“模拟任务 #5”。"],
        ["ParticleDetection", "当前共有 41 条颗粒明细，挂载在上述 5 条检测记录下。"],
        ["ProcessParameterHistory", "当前为空；代码中的模拟数据不会写入该表。"],
    ],
    [2.3, 7.2],
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("MzTLZ 数据库表结构说明")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(85, 85, 85)

doc.save(OUT)
print(OUT)
